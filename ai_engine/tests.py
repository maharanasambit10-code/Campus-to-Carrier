from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from docx import Document

from accounts.models import Membership, User
from companies.models import Company
from jobs.models import Job
from students.models import Skill, StudentProfile, StudentSkill
from .models import ResumeAnalysis
from .services import match_jobs


class CareerMatchTests(TestCase):
	def setUp(self):
		self.user = User.objects.create_user(username='student', password='password')
		Membership.objects.create(user=self.user, plan='PRO_MONTHLY', status='ACTIVE')
		self.profile = StudentProfile.objects.create(user=self.user, degree='B.Tech', cgpa=8.0)
		for name in ('Python', 'Django', 'SQL'):
			skill = Skill.objects.create(name=name)
			StudentSkill.objects.create(student=self.profile, skill=skill)
		Skill.objects.create(name='React')
		company = Company.objects.create(name='Campus Tech')
		self.job = Job.objects.create(
			company=company,
			title='Python Developer',
			description='Backend role',
			location='Remote',
			graduation_year=2026,
			application_deadline='2027-01-01T00:00:00Z',
		)
		self.job.required_skills.set(Skill.objects.filter(name__in=['Python', 'Django', 'SQL', 'React']))

	def test_partial_match_exposes_missing_skills(self):
		match = match_jobs(self.profile, [self.job])[0]
		self.assertEqual(match['score'], 75)
		self.assertEqual(match['status'], 'PARTIAL MATCH')
		self.assertEqual(match['missing_skills'], ['react'])

	def test_docx_upload_creates_structured_analysis(self):
		document = Document()
		document.add_paragraph('Asha Student')
		document.add_paragraph('B.Tech Computer Science 2026')
		document.add_paragraph('Python Django SQL Git')
		content = BytesIO()
		document.save(content)
		upload = SimpleUploadedFile('resume.docx', content.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

		self.client.force_login(self.user)
		response = self.client.post('/student/career-match/', {'resume': upload})

		self.assertEqual(response.status_code, 302)
		analysis = ResumeAnalysis.objects.get(student=self.profile)
		self.assertIn('Python', analysis.skills)
		self.assertFalse(analysis.parse_error)
		page = self.client.get('/student/career-match/')
		self.assertEqual(page.status_code, 200)
	def test_trust_layer_match_explanation(self):
		match = match_jobs(self.profile, [self.job])[0]
		self.assertIn('trust_layer', match)
		trust = match['trust_layer']
		self.assertIn('evidence_behind_the_role', trust)
		self.assertIn('demonstrated_strengths', trust)
		self.assertIn('inferred_skills', trust)
		self.assertIn('next_practice_mission', trust)
		self.assertEqual(trust['flight_challenge_slug'], 'high-concurrency-api')


class FlightSimulatorAndPassportTests(TestCase):
	def setUp(self):
		self.user = User.objects.create_user(username='student_user', password='password123')
		self.profile = StudentProfile.objects.create(user=self.user, degree='B.Tech CSE', cgpa=8.5)
		self.client.force_login(self.user)

	def test_simulator_catalog_and_seeding(self):
		response = self.client.get('/simulator/')
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Career Flight Simulator')
		self.assertContains(response, 'Customer Retention &amp; Churn')

	def test_simulator_submission_and_evaluation(self):
		from .services import seed_default_flight_challenges
		seed_default_flight_challenges()

		response = self.client.post('/simulator/data-storytelling-churn/submit/', {
			'solution_text': 'Our cohort retention telemetry shows user drop-off in cohort B due to delivery friction and auto-debit confusion. Metric drop is 14%.',
			'recommendations': '1. Pilot proactive delivery notifications. 2. Implement a 30-day renewal confirmation email. 3. Monitor retention KPI weekly.',
			'artifact_url': 'https://github.com/example/retention-story',
		})
		self.assertEqual(response.status_code, 302)

		from .models import SimulationSubmission, ProofPassport
		submission = SimulationSubmission.objects.filter(student=self.profile).first()
		self.assertIsNotNone(submission)
		self.assertGreaterEqual(submission.overall_score, 70)
		self.assertTrue(len(submission.demonstrated_strengths) > 0)
		self.assertTrue(len(submission.skill_gaps) > 0)
		self.assertIn('slug', submission.next_practice_mission)

		# Verify automatic Proof Passport synchronization
		passport = ProofPassport.objects.filter(student=self.profile).first()
		self.assertIsNotNone(passport)
		sim_artifact = passport.artifacts.filter(artifact_type='SIMULATION').first()
		self.assertIsNotNone(sim_artifact)
		self.assertEqual(sim_artifact.confidence_level, 'DEMONSTRATED')

	def test_opportunity_compiler_flow(self):
		company = Company.objects.create(name='Acme Corp')
		job = Job.objects.create(
			company=company,
			title='Full-Stack Engineer',
			description='Must have experience with asynchronous python, docker, rest api, and react.',
			location='Remote',
			application_deadline='2027-01-01T00:00:00Z',
		)

		response = self.client.post('/compiler/run/', {
			'job_id': str(job.id),
			'schedule_type': '2_WEEK_SPRINT',
			'available_hours': '12',
		})
		self.assertEqual(response.status_code, 302)

		from .models import OpportunityCompilerSession
		session = OpportunityCompilerSession.objects.filter(student=self.profile).first()
		self.assertIsNotNone(session)
		self.assertTrue(len(session.plain_language_skills) > 0)
		self.assertTrue(len(session.evidence_mapping) > 0)
		self.assertTrue(len(session.sprint_plan) > 0)

		# Test detail view
		detail = self.client.get(f'/compiler/session/{session.id}/')
		self.assertEqual(detail.status_code, 200)
		self.assertContains(detail, 'Requirements → Plain-Language Skills')

		# Test toggle milestone
		toggle = self.client.post(f'/compiler/session/{session.id}/task/0/toggle/')
		self.assertEqual(toggle.status_code, 200)

		# Test public proof brief
		brief = self.client.get(f'/compiler/proof-brief/{session.proof_brief_token}/')
		self.assertEqual(brief.status_code, 200)
		self.assertContains(brief, 'ROLE-READY PROOF BRIEF')

	def test_proof_passport_dashboard_and_public_view(self):
		response = self.client.get('/passport/')
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Evidence, Not Empty Claims')

		from .models import ProofPassport
		passport = ProofPassport.objects.get(student=self.profile)

		# Test public verifiable view (accessible without login)
		anon_client = self.client_class()
		pub = anon_client.get(f'/passport/view/{passport.public_share_token}/')
		self.assertEqual(pub.status_code, 200)
		self.assertContains(pub, 'Verified Candidate Proof Passport')

		# Test add custom artifact
		add_art = self.client.post('/passport/artifact/add/', {
			'title': 'High Throughput Message Bus',
			'artifact_type': 'PROJECT',
			'description': 'Built event queue with Redis and Python.',
			'skills_evidenced': 'Python, Redis, Queues',
			'external_url': 'https://github.com/example/bus',
			'score_or_grade': 'Tested',
		})
		self.assertEqual(add_art.status_code, 302)
		self.assertTrue(passport.artifacts.filter(title='High Throughput Message Bus').exists())

		# Test add peer review
		add_peer = self.client.post('/passport/peer-review/add/', {
			'reviewer_name': 'Sarah Connor',
			'reviewer_role': 'Tech Lead',
			'project_name': 'Core Platform',
			'review_text': 'Excellent distributed systems thinking and high code quality.',
			'skills_endorsed': 'Python, Architecture',
		})
		self.assertEqual(add_peer.status_code, 302)
		self.assertTrue(self.profile.peer_reviews.filter(reviewer_name='Sarah Connor').exists())
